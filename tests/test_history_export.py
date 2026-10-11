from __future__ import annotations

import unittest
from io import BytesIO

from openpyxl import load_workbook

from parsing_site.history_export import build_history_workbook


class HistoryExportTest(unittest.TestCase):
    def test_builds_filterable_workbook_with_typed_values(self) -> None:
        content = build_history_workbook(
            [
                {
                    "ts": "2026-10-08T12:34:56+03:00",
                    "static_index": 3,
                    "static_name": "Пост наблюдения",
                    "static_lat": 47.222,
                    "static_lng": 39.718,
                    "dynamic_index": 17,
                    "dynamic_name": "Машина",
                    "dynamic_lat": 47.223,
                    "dynamic_lng": 39.719,
                    "distance": 0.145,
                    "status": "alert",
                }
            ],
            date_from="2026-10-01",
            date_to="2026-10-08",
            static_index=3,
            max_distance_km=10,
            dynamic_query="Маш",
            source="real",
        )

        workbook = load_workbook(BytesIO(content))
        sheet = workbook["Хронология"]
        self.assertEqual(sheet["A2"].value, "Хронология сближений")
        self.assertEqual(sheet["E3"].value, "#3 Пост наблюдения")
        self.assertEqual(sheet.freeze_panes, "A9")
        self.assertEqual(sheet["A8"].value, "Дата и время")
        self.assertEqual(sheet["A9"].value.year, 2026)
        self.assertIsInstance(sheet["D9"].value, float)
        self.assertEqual(sheet["J9"].value, 0.145)
        self.assertEqual(sheet["K9"].value, "Красная зона")
        self.assertIn("ChronologyTable", sheet.tables)

    def test_empty_export_still_contains_headers_and_filter(self) -> None:
        content = build_history_workbook(
            [],
            date_from="2026-10-01",
            date_to="2026-10-08",
            static_index=1,
            max_distance_km=5,
            dynamic_query="",
            source="fake",
        )

        sheet = load_workbook(BytesIO(content))["Хронология"]
        self.assertEqual(sheet["E5"].value, 0)
        self.assertEqual(sheet.auto_filter.ref, "A8:K8")
        self.assertEqual(len(sheet.tables), 0)


if __name__ == "__main__":
    unittest.main()
